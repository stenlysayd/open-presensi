import io
from datetime import date
from typing import List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_attendance_excel(
    organization_name: str,
    start_date: date,
    end_date: date,
    records: List[Any]
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Presensi"

    # Title Banner
    ws.merge_cells("A1:G1")
    ws["A1"] = f"REKAP PRESENSI - {organization_name.upper()}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Periode: {start_date.strftime('%d/%m/%Y')} s/d {end_date.strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    # Table Header
    headers = ["No", "No. Induk / Reg", "Nama Anggota", "Grup / Kelas", "Tipe Absen", "Waktu", "Status"]
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Rows
    current_row = 5
    for idx, r in enumerate(records, 1):
        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=r.get("reg_number", "-")).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=3, value=r.get("name", "-"))
        ws.cell(row=current_row, column=4, value=r.get("group_name", "-"))
        ws.cell(row=current_row, column=5, value=r.get("record_type", "-").upper()).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=6, value=r.get("time_str", "-")).alignment = Alignment(horizontal="center")
        
        status_cell = ws.cell(row=current_row, column=7, value=r.get("status", "-").upper())
        status_cell.alignment = Alignment(horizontal="center")
        
        # Color badge for status
        if r.get("status") == "hadir":
            status_cell.font = Font(color="27AE60", bold=True)
        elif r.get("status") in ["izin", "sakit"]:
            status_cell.font = Font(color="D35400", bold=True)
        else:
            status_cell.font = Font(color="C0392B", bold=True)

        for c in range(1, 8):
            ws.cell(row=current_row, column=c).border = thin_border

        current_row += 1

    # Adjust Column Widths
    col_widths = [6, 18, 28, 20, 14, 18, 14]
    for i, col_letter in enumerate(["A", "B", "C", "D", "E", "F", "G"]):
        ws.column_dimensions[col_letter].width = col_widths[i]

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def generate_attendance_pdf(
    organization_name: str,
    start_date: date,
    end_date: date,
    records: List[Any]
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('#1F497D'),
        alignment=1, # Center
        spaceAfter=5
    )
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.gray,
        alignment=1,
        spaceAfter=15
    )
    elements.append(Paragraph(f"LAPORAN REKAP PRESENSI - {organization_name.upper()}", title_style))
    elements.append(Paragraph(f"Periode: {start_date.strftime('%d/%m/%Y')} s/d {end_date.strftime('%d/%m/%Y')}", subtitle_style))

    # Table Data
    table_data = [["No", "No. Induk", "Nama Anggota", "Grup", "Tipe", "Waktu", "Status"]]
    for idx, r in enumerate(records[:150], 1): # Limit for single PDF demo
        table_data.append([
            str(idx),
            r.get("reg_number", "-"),
            r.get("name", "-"),
            r.get("group_name", "-"),
            r.get("record_type", "").upper(),
            r.get("time_str", "-"),
            r.get("status", "").upper()
        ])

    pdf_table = Table(table_data, colWidths=[25, 65, 140, 90, 50, 95, 60])
    pdf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F497D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'LEFT'), # Nama left aligned
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')),
    ]))

    elements.append(pdf_table)
    doc.build(elements)
    return buffer.getvalue()
